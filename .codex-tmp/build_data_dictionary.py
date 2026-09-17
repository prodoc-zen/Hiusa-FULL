from __future__ import annotations

import re
import sqlite3
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_DB = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
OLD_DICTIONARY = ROOT / "docs2" / "Hiusa_Data_Dictionary_UPDATED.docx"
MIGRATIONS = ROOT / "server" / "database" / "migrations"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
HEADER_FILL = "E8EEF5"
ALT_FILL = "F8FAFC"
BORDER = "7F8C9A"
WHITE = "FFFFFF"

TABLE_ORDER = [
    "organizations",
    "academic_programs",
    "academic_sections",
    "users",
    "sbo_positions",
    "fingerprints",
    "fingerprint_verifications",
    "password_reset_tokens",
    "sessions",
    "personal_access_tokens",
    "announcements",
    "announcement_recipients",
    "announcement_views",
    "notifications",
    "events",
    "tasks",
    "task_progress_updates",
    "ai_outputs",
    "task_recommendations",
    "budgets",
    "transactions",
    "financial_forecasts",
    "financial_reports",
    "financial_report_deadlines",
    "collections",
    "remittances",
    "cash_advances",
    "cash_advance_repayments",
    "invoices",
    "invoice_payments",
    "merchandise",
    "orders",
    "approval_requests",
    "audit_logs",
    "elections",
    "election_positions",
    "partylists",
    "candidates",
    "votes",
    "attendance",
    "cache",
    "cache_locks",
    "jobs",
    "job_batches",
    "failed_jobs",
]

OLD_TABLE_ORDER = [
    "organizations", "users", "password_reset_tokens", "sessions",
    "announcements", "events", "tasks", "budgets", "financial_forecasts",
    "merchandise", "orders", "transactions", "elections",
    "election_positions", "candidates", "votes", "notifications",
    "attendance", "partylists", "personal_access_tokens", "cache",
    "cache_locks", "jobs", "job_batches", "failed_jobs",
]

TABLE_DESCRIPTIONS = {
    "organizations": "Stores organization workspaces, including the Student Affairs Office and student organizations. It is the principal ownership boundary used to isolate records and users.",
    "academic_programs": "Defines the academic programs available within an organization for student profile classification.",
    "academic_sections": "Defines year-level and section combinations under an academic program.",
    "users": "Stores authenticated accounts for the SAO/Super Admin, organization administrators, SBO officers, department heads, and students. The school_id field is the primary key used by all user relationships.",
    "sbo_positions": "Defines organization positions and the system role associated with each position title.",
    "fingerprints": "Stores organization-scoped biometric templates enrolled for attendance verification. Templates are retained as long text with an explicit template format.",
    "fingerprint_verifications": "Records biometric verification attempts, scores, thresholds, operators, and optional event context for accountability.",
    "password_reset_tokens": "Stores temporary password-recovery tokens, scoped by organization and email address.",
    "sessions": "Stores Laravel database-session payloads and activity metadata when database sessions are enabled.",
    "personal_access_tokens": "Stores Laravel Sanctum API tokens used for authenticated API access.",
    "announcements": "Stores organization and SAO announcements, publication workflow state, audience targeting, scheduling, media, and priority flags.",
    "announcement_recipients": "Materializes the users selected to receive a targeted announcement.",
    "announcement_views": "Records the first view of an announcement by each user for reach and engagement reporting.",
    "notifications": "Stores user-specific notifications, delivery scheduling, read state, and links to related records.",
    "events": "Stores organization activities, planning details, approval state, schedule, location, media, and budget requirements.",
    "tasks": "Stores event and organization work items, assignments, dependencies, priorities, progress, and AI-assisted delegation results.",
    "task_progress_updates": "Stores the chronological progress history and notes submitted for a task.",
    "ai_outputs": "Stores generated AI decision-support output, structured inputs and outputs, version history, errors, and human decision state.",
    "task_recommendations": "Stores ranked AI-supported officer recommendations and the scoring factors used for task delegation.",
    "budgets": "Stores organization and event budget allocations together with remaining balances and AI-generated spending guidance.",
    "transactions": "Stores the organization financial ledger for income and expense records, including receipts, events, payers, and budgets.",
    "financial_forecasts": "Stores generated financial projections and safe-spending estimates for a reporting period.",
    "financial_reports": "Stores generated financial reports, covered periods, ledger sources, exports, signatories, supporting documents, submission state, and two-stage approval data.",
    "financial_report_deadlines": "Stores SAO-defined financial-report submission deadlines and the announcement generated for each deadline.",
    "collections": "Stores collected funds awaiting verification and optional posting to the financial ledger.",
    "remittances": "Stores transfers of verified collections and the users responsible for remitting and verifying them.",
    "cash_advances": "Stores requested, approved, released, and repaid cash advances tied to organizations and optional events.",
    "cash_advance_repayments": "Stores repayment installments for cash advances and their optional ledger transactions.",
    "invoices": "Stores student financial obligations related to events or merchandise orders.",
    "invoice_payments": "Stores payments applied to invoices and their optional ledger transactions.",
    "merchandise": "Stores organization merchandise, prices, inventory quantities, media, and availability state.",
    "orders": "Stores student merchandise orders, payment evidence, multi-stage review, claim verification, and linked ledger transactions.",
    "approval_requests": "Stores approval workflow requests for related entities, required roles, assigned approvers, decisions, remarks, and audit timestamps.",
    "audit_logs": "Stores immutable accountability records describing who performed an action and how a record changed.",
    "elections": "Stores election schedules, lifecycle state, approval, result visibility, media, and organization ownership.",
    "election_positions": "Defines the positions available in an election and the maximum number of winners for each position.",
    "partylists": "Stores organization-scoped election groups and their identity, description, and banner media.",
    "candidates": "Stores approved election candidates, their positions, platforms, media, and optional partylist membership.",
    "votes": "Stores cast ballots with integrity hashes and uniqueness controls that prevent duplicate votes per position.",
    "attendance": "Stores event check-in and check-out records, attendance method, status, operator, and remarks.",
    "cache": "Stores Laravel application cache entries and expiration values.",
    "cache_locks": "Stores Laravel atomic cache locks used to coordinate concurrent operations.",
    "jobs": "Stores queued Laravel jobs awaiting or undergoing processing.",
    "job_batches": "Stores aggregate state for batches of queued Laravel jobs.",
    "failed_jobs": "Stores queued jobs that exhausted their attempts, including payload and exception details.",
}

FIELD_DESCRIPTIONS = {
    "school_id": "Primary school identifier for the user account; referenced instead of a users.id column.",
    "organization_type": "Classifies the workspace as a student organization or system-administration office.",
    "announcement_source": "Identifies whether the announcement originated from an organization or the SAO.",
    "target_scope": "Defines the audience-selection strategy for the announcement.",
    "target_organization_ids": "JSON array of organizations targeted by the announcement.",
    "target_departments": "JSON array of academic departments targeted by the announcement.",
    "target_roles": "JSON array of user roles targeted by the announcement.",
    "signatories": "JSON snapshot of the required Treasurer, President, Adviser, and SBO Adviser signatories.",
    "supporting_documents": "JSON list of supporting-document metadata attached to the report.",
    "submission_status": "Current financial-report workflow state from draft through final SAO approval or rejection.",
    "deadline_at": "Date and time by which organizations must submit their financial reports.",
    "deadline_id": "References financial_report_deadlines(id) applied when the report was submitted.",
    "department_head_approved_by": "References users(school_id) for the Department Head who completed first-stage approval.",
    "department_head_approved_at": "Date and time of Department Head approval.",
    "sao_approved_by": "References users(school_id) for the SAO/Super Admin who completed final approval.",
    "sao_approved_at": "Date and time of final SAO approval.",
    "biometric_template": "Legacy DigitalPersona-compatible binary biometric template stored on the user record.",
    "template": "Encoded biometric template used for fingerprint matching.",
    "template_format": "Format and version identifier used to interpret the biometric template.",
    "finger_index": "Numeric identifier for the enrolled finger.",
    "score": "Similarity score produced by biometric verification.",
    "threshold": "Minimum verification score required for a successful match.",
    "active_key": "Unique key that prevents duplicate active approval requests for the same entity.",
    "decision": "Normalized approval decision recorded by the approver.",
    "source_transaction_ids": "Serialized list of ledger transaction identifiers included in the report.",
    "structured_input": "JSON snapshot of the normalized input supplied to the AI workflow.",
    "structured_output": "JSON representation of the AI workflow result.",
    "context_version": "Version identifier for the prompt and decision-support context.",
    "decision_status": "Human review state for the generated AI output.",
    "delegation_snapshot": "JSON snapshot of delegation scores and context used when the task was assigned.",
    "weights": "JSON object containing the scoring weights used by the recommendation engine.",
    "vote_hash": "Unique integrity value used to detect duplicate or altered ballots.",
    "payment_reference_key": "Normalized unique payment reference used to prevent duplicate order payment records.",
    "receipt_number": "Event-scoped receipt sequence number.",
}


def set_font(run, size=None, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[i] / 1440)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run("Page ")
    set_font(run, size=8.5, color=MUTED)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def load_existing_descriptions():
    result = {}
    if not OLD_DICTIONARY.exists():
        return result
    old = Document(OLD_DICTIONARY)
    for table_name, table in zip(OLD_TABLE_ORDER, old.tables):
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) == 4 and cells[0]:
                result[(table_name, cells[0])] = cells[3]
    return result


def migration_type_hints():
    hints = {}
    call_re = re.compile(r"\$table->(\w+)\(\s*'([^']+)'(?:\s*,\s*(.*?))?\)\s*(?:->|;)")
    table_re = re.compile(r"Schema::(?:create|table)\(\s*'([^']+)'")
    for path in sorted(MIGRATIONS.glob("*.php")):
        source = path.read_text(encoding="utf-8")
        source = source.split("public function down", 1)[0]
        current = None
        for line in source.splitlines():
            match = table_re.search(line)
            if match:
                current = match.group(1)
                hints.setdefault(current, {})
                continue
            if current and "$table->id();" in line:
                hints[current]["id"] = "bigint unsigned"
            if current and "$table->timestamps();" in line:
                hints[current]["created_at"] = "timestamp"
                hints[current]["updated_at"] = "timestamp"
            if current:
                call = call_re.search(line)
                if call:
                    method, field, args = call.group(1), call.group(2), call.group(3) or ""
                    hints[current][field] = laravel_type(method, args)
            if current and line.strip() == "});":
                current = None
    return hints


def laravel_type(method, args):
    args = args.strip()
    if method in {"string", "char"}:
        match = re.search(r"(\d+)", args)
        length = match.group(1) if match else "255"
        return f"varchar({length})" if method == "string" else f"char({length})"
    mapping = {
        "id": "bigint unsigned",
        "foreignId": "bigint unsigned",
        "unsignedBigInteger": "bigint unsigned",
        "bigInteger": "bigint",
        "unsignedInteger": "int unsigned",
        "integer": "int",
        "unsignedSmallInteger": "smallint unsigned",
        "smallInteger": "smallint",
        "unsignedTinyInteger": "tinyint unsigned",
        "tinyInteger": "tinyint",
        "boolean": "tinyint(1)",
        "text": "text",
        "mediumText": "mediumtext",
        "longText": "longtext",
        "json": "json",
        "binary": "blob",
        "date": "date",
        "dateTime": "datetime",
        "timestamp": "timestamp",
        "time": "time",
        "enum": "enum",
        "uuid": "char(36)",
    }
    if method == "decimal":
        nums = re.findall(r"\d+", args)
        return f"decimal({nums[0]},{nums[1]})" if len(nums) >= 2 else "decimal(8,2)"
    return mapping.get(method, method)


def schema_snapshot(connection):
    tables = {}
    table_names = [
        row[0]
        for row in connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table' "
            "AND name NOT IN ('migrations', 'sqlite_sequence') ORDER BY name"
        )
    ]
    for table_name in table_names:
        columns = connection.execute(f'PRAGMA table_info("{table_name}")').fetchall()
        foreign_keys = connection.execute(f'PRAGMA foreign_key_list("{table_name}")').fetchall()
        indexes = connection.execute(f'PRAGMA index_list("{table_name}")').fetchall()
        fk_by_field = {
            row[3]: {"table": row[2], "column": row[4], "on_delete": row[6]}
            for row in foreign_keys
        }
        index_data = []
        for index in indexes:
            index_name = index[1]
            fields = [row[2] for row in connection.execute(f'PRAGMA index_info("{index_name}")')]
            index_data.append({"name": index_name, "unique": bool(index[2]), "fields": fields})
        tables[table_name] = {
            "columns": [
                {
                    "name": row[1],
                    "sqlite_type": row[2],
                    "not_null": bool(row[3]) or bool(row[5]),
                    "default": row[4],
                    "pk": bool(row[5]),
                }
                for row in columns
            ],
            "foreign_keys": fk_by_field,
            "indexes": index_data,
        }
    return tables


def data_type(table_name, column, fk, hints, old_descriptions):
    name = column["name"]
    if name == "school_id":
        return "int unsigned"
    if table_name == "users" and name == "biometric_template":
        return "longblob"
    if table_name in hints and name in hints[table_name]:
        return hints[table_name][name]
    if fk:
        return "int unsigned" if fk["table"] == "users" else "bigint unsigned"
    sqlite_type = column["sqlite_type"].lower()
    if sqlite_type.startswith("tinyint"):
        return "tinyint(1)"
    if "int" in sqlite_type:
        return "bigint unsigned" if name == "id" else "int"
    if sqlite_type == "numeric":
        if name in {"score", "threshold"}:
            return "decimal(10,4)"
        if name.endswith("_score") or name == "total_score":
            return "decimal(6,2)"
        return "decimal(12,2)"
    if sqlite_type == "text":
        if name in {
            "structured_input", "structured_output", "delegation_snapshot", "weights",
            "notification_preferences", "model_details", "planning_details", "signatories",
            "supporting_documents", "target_organization_ids", "target_departments", "target_roles",
        }:
            return "json"
        return "text"
    if sqlite_type == "blob":
        return "longblob"
    if sqlite_type == "datetime":
        return "timestamp"
    if sqlite_type == "date":
        return "date"
    if sqlite_type.startswith("varchar"):
        return "varchar(255)"
    return column["sqlite_type"] or "varchar(255)"


def format_default(value):
    if value is None:
        return None
    value = str(value)
    if value.upper() == "CURRENT_TIMESTAMP":
        return "CURRENT_TIMESTAMP"
    return value.strip("'")


def constraints_for(table_name, column, table_data):
    name = column["name"]
    parts = []
    if column["pk"]:
        parts.append("PK")
    fk = table_data["foreign_keys"].get(name)
    if fk:
        parts.append(f"FK -> {fk['table']}({fk['column']})")
    parts.append("NOT NULL" if column["not_null"] else "NULL")
    if column["pk"] and name == "id" and "int" in column["sqlite_type"].lower():
        parts.append("AUTO_INCREMENT")
    default = format_default(column["default"])
    if default is not None:
        parts.append(f"DEFAULT {default}")
    for index in table_data["indexes"]:
        fields = index["fields"]
        if not fields or fields[0] != name or index["name"].startswith("sqlite_autoindex"):
            continue
        label = "UNIQUE" if index["unique"] else "INDEX"
        if len(fields) > 1:
            label += " (" + ", ".join(fields) + ")"
        parts.append(label)
    return ", ".join(parts)


def humanize(value):
    return value.replace("_", " ").strip()


def singular(value):
    if value.endswith("ies"):
        return value[:-3] + "y"
    if value.endswith("sses"):
        return value[:-2]
    if value.endswith("s") and not value.endswith("ss"):
        return value[:-1]
    return value


def description_for(table_name, field, table_data, existing):
    if field in FIELD_DESCRIPTIONS:
        return FIELD_DESCRIPTIONS[field]
    if (table_name, field) in existing:
        return existing[(table_name, field)]
    fk = table_data["foreign_keys"].get(field)
    if fk:
        related = humanize(singular(fk["table"]))
        return f"References {fk['table']}({fk['column']}) and links this record to the related {related}."
    if field == "id":
        return f"Unique identifier for the {humanize(singular(table_name))} record."
    if field == "created_at":
        return "Date and time when the record was created."
    if field == "updated_at":
        return "Date and time when the record was last updated."
    if field == "status":
        return "Current lifecycle or workflow status of the record."
    if field == "name":
        return f"Display name of the {humanize(singular(table_name))}."
    if field == "title":
        return "Short title displayed for the record."
    if field == "description":
        return "Detailed description of the record."
    if field == "notes":
        return "Optional administrative notes about the record."
    if field == "remarks" or field.endswith("_remarks"):
        return "Optional remarks recorded during processing or review."
    if field == "role" or field.endswith("_role"):
        return "System role used for authorization or workflow routing."
    if field.endswith("_at"):
        return f"Date and time when {humanize(field[:-3])} occurred."
    if field.endswith("_date"):
        return f"Calendar date for {humanize(field[:-5])}."
    if field.endswith("_url"):
        return f"Stored URL for the related {humanize(field[:-4])} resource."
    if field.startswith("is_") or field in {"results_visible", "requires_budget"}:
        return f"Boolean flag indicating whether {humanize(field.removeprefix('is_'))} applies."
    if field.endswith("_amount") or field in {"amount", "price", "total_price", "predicted_income", "predicted_expense", "predicted_balance", "safe_spending_limit", "remaining_amount", "recommended_allocation"}:
        return f"Monetary value for {humanize(field)}."
    if field.endswith("_count") or field in {"quantity", "stock_quantity", "attempts", "rank", "sequence", "progress_percent", "year_level", "max_winners"}:
        return f"Numeric value representing {humanize(field)}."
    if field.endswith("_text") or field in {"body", "purpose", "platform", "payload", "exception", "error_message", "prompt_text", "output_text"}:
        return f"Text content for {humanize(field)}."
    if field.endswith("_type") or field == "type":
        return f"Classification value for {humanize(field)}."
    if field.endswith("_reference") or field == "reference":
        return f"Reference value used to identify or reconcile the {humanize(singular(table_name))}."
    if field in {"email", "contact_number", "ip_address", "location", "category", "source", "method", "priority", "phase"}:
        return f"Stored {humanize(field)} value for the record."
    return f"Stores the {humanize(field)} value for this {humanize(singular(table_name))} record."


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 8),
        ("Subtitle", 13, MUTED, 0, 18),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_running_furniture(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("HIUSA DATA DICTIONARY  |  CURRENT SCHEMA")
    set_font(run, size=8, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    add_page_field(footer)


def add_cover(doc, table_count, field_count):
    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(8)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_font(kicker.add_run("DATABASE SCHEMA REFERENCE"), size=10, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("HIUSA Data Dictionary")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Current Laravel / MySQL Schema")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(26)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(10)
    set_font(summary.add_run(f"{table_count} tables  |  {field_count} fields"), size=12, bold=True, color=DARK_BLUE)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    set_font(date.add_run("Schema snapshot: September 16, 2026"), size=10.5, color=MUTED)
    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(0)
    set_font(source.add_run("Derived from a clean execution of all repository migrations."), size=9.5, italic=True, color=MUTED)


def add_overview(doc, schema):
    doc.add_page_break()
    doc.add_heading("About this dictionary", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "This document replaces the earlier 25-table reference with the complete current schema. "
        "It includes organization isolation, SAO governance, approval routing, financial reporting, "
        "AI auditability, academic structure, biometric attendance, and Laravel framework tables."
    )
    doc.add_heading("Constraint notation", level=2)
    legend = doc.add_table(rows=1, cols=2)
    legend.style = "Table Grid"
    legend.rows[0].cells[0].text = "Notation"
    legend.rows[0].cells[1].text = "Meaning"
    legend_items = [
        ("PK", "Primary key"),
        ("FK", "Foreign key and referenced table/field"),
        ("NOT NULL / NULL", "Whether a value is required"),
        ("UNIQUE", "Uniqueness constraint or composite unique key"),
        ("INDEX", "Lookup or composite performance index"),
        ("DEFAULT", "Database default value"),
    ]
    for label, meaning in legend_items:
        cells = legend.add_row().cells
        cells[0].text = label
        cells[1].text = meaning
    set_table_geometry(legend, [2100, 7260])
    set_repeat_table_header(legend.rows[0])
    for i, row in enumerate(legend.rows):
        prevent_row_split(row)
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, HEADER_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_font(run, size=9, bold=(i == 0), color=DARK_BLUE if i == 0 else INK)
    doc.add_heading("Table inventory", level=2)
    inventory = doc.add_table(rows=0, cols=2)
    inventory.style = "Table Grid"
    half = (len(TABLE_ORDER) + 1) // 2
    for index in range(half):
        cells = inventory.add_row().cells
        left_number = index + 1
        cells[0].text = f"{left_number}. {TABLE_ORDER[index]}"
        right_index = index + half
        if right_index < len(TABLE_ORDER):
            cells[1].text = f"{right_index + 1}. {TABLE_ORDER[right_index]}"
    set_table_geometry(inventory, [4680, 4680])
    for row_index, row in enumerate(inventory.rows):
        prevent_row_split(row)
        for cell in row.cells:
            set_cell_margins(cell, top=55, bottom=55)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, ALT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, size=8.5, color=INK)


def add_dictionary_table(doc, number, table_name, table_data, type_hints, existing):
    doc.add_page_break()
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(2)
    caption.paragraph_format.keep_with_next = True
    set_font(caption.add_run(f"Table {number}"), size=9.5, color=INK)
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True
    set_font(heading.add_run(f"DATA DICTIONARY OF {table_name.upper()} TABLE"), size=10.5, bold=True, color=DARK_BLUE)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Field Name", "Constraints", "Data Type", "Description"]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for column in table_data["columns"]:
        row = table.add_row()
        field = column["name"]
        fk = table_data["foreign_keys"].get(field)
        values = [
            field,
            constraints_for(table_name, column, table_data),
            data_type(table_name, column, fk, type_hints, existing),
            description_for(table_name, field, table_data, existing),
        ]
        for cell, value in zip(row.cells, values):
            cell.text = value
    widths = [1800, 2550, 1500, 3510]
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for col_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif row_index % 2 == 0:
                set_cell_shading(cell, ALT_FILL)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if col_index in (0, 2):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_font(
                        run,
                        size=8.2 if row_index else 8.5,
                        bold=(row_index == 0),
                        color=DARK_BLUE if row_index == 0 else INK,
                    )

    narrative = doc.add_paragraph()
    narrative.paragraph_format.space_before = Pt(9)
    narrative.paragraph_format.space_after = Pt(4)
    narrative.paragraph_format.line_spacing = 1.2
    narrative.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    narrative.add_run(TABLE_DESCRIPTIONS[table_name])
    foreign_count = len(table_data["foreign_keys"])
    unique_count = sum(1 for index in table_data["indexes"] if index["unique"] and not index["name"].startswith("sqlite_autoindex"))
    controls = doc.add_paragraph()
    controls.paragraph_format.space_before = Pt(0)
    controls.paragraph_format.space_after = Pt(0)
    controls.paragraph_format.line_spacing = 1.2
    controls.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    relation_text = (
        f"The final migrated definition contains {len(table_data['columns'])} fields, "
        f"{foreign_count} foreign-key relationship{'s' if foreign_count != 1 else ''}, and "
        f"{unique_count} explicit unique constraint{'s' if unique_count != 1 else ''}. "
        "Constraints shown above reflect the current forward migration state."
    )
    set_font(controls.add_run(relation_text), size=9.5, color=MUTED)


def build_document():
    connection = sqlite3.connect(SCHEMA_DB)
    schema = schema_snapshot(connection)
    connection.close()
    missing = [name for name in TABLE_ORDER if name not in schema]
    extra = [name for name in schema if name not in TABLE_ORDER]
    if missing or extra:
        raise RuntimeError(f"Schema inventory mismatch. Missing={missing}; Extra={extra}")
    existing = load_existing_descriptions()
    type_hints = migration_type_hints()

    doc = Document()
    doc.core_properties.title = "HIUSA Data Dictionary"
    doc.core_properties.subject = "Current Laravel and MySQL database schema reference"
    doc.core_properties.author = "HIUSA Development Team"
    doc.core_properties.keywords = "HIUSA, data dictionary, database, Laravel, MySQL"
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    add_running_furniture(section)
    total_fields = sum(len(schema[name]["columns"]) for name in TABLE_ORDER)
    add_cover(doc, len(TABLE_ORDER), total_fields)
    add_overview(doc, schema)
    for number, table_name in enumerate(TABLE_ORDER, start=1):
        add_dictionary_table(doc, number, table_name, schema[table_name], type_hints, existing)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Tables: {len(TABLE_ORDER)}")
    print(f"Fields: {total_fields}")


if __name__ == "__main__":
    build_document()
